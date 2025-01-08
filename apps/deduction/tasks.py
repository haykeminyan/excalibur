# tasks.py (inside your app, e.g., "deduction")
from celery import shared_task
from django.http import JsonResponse
from docx import Document
from django.core.files.storage import default_storage
from django.conf import settings
import os, io, re
from celery.result import AsyncResult
from excalibur.rabbitmq_service import send_rabbitmq_message
from .models import Deduction
from .constants import DEDUCTION_FIELDS
from apps.facture.parsing_docx import replace_placeholders_in_doc, set_font_size
from django.forms.models import model_to_dict
from django.urls import reverse
import logging


logger = logging.getLogger(__name__)


@shared_task
def generate_docx_task(deduction_id):
    # Get the Deduction object based on the ID
    deduction_object = Deduction.objects.get(pk=deduction_id)

    # Generate the docx file
    facture_dict = model_to_dict(deduction_object)
    template_path = os.path.join(settings.BASE_DIR, 'apps/deduction/static/deduction/file_templates/file_input/Deduction_template.docx')
    doc = Document(template_path)

    for field in DEDUCTION_FIELDS:
        regex = re.compile(rf'{re.escape(field)}')
        replace_placeholders_in_doc(doc, regex, facture_dict)

    set_font_size(doc)

    # Create a file-like object to store the generated doc
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    # Create the path where the file will be saved
    file_path = f'/usr/src/app/apps/deduction/static/deduction/file_templates/file_output/deduction_{deduction_object.number_deduction}.docx'
    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    with open(file_path, 'wb') as docx_file:
        docx_file.write(file_stream.getvalue())

    # Optionally, you can send a message to RabbitMQ here if needed.
    message = f"Deduction document exported: ID={deduction_object.id}, Owner={deduction_object.owner}"
    send_rabbitmq_message(Deduction, message)

    return file_path


def task_status(request, task_id):
    task = AsyncResult(task_id)
    logger.info(task.state)
    logger.info('!'*100)
    # Check if the task is complete
    if task.state == 'SUCCESS':
        # The file is ready, return the download URL
        file_path = task.result  # This is the file path returned by the task
        file_url = reverse('apps.deduction:download_deduction_document', kwargs={'file_name': os.path.basename(file_path)})
        return JsonResponse({'status': 'Completed', 'file_url': file_url})

    elif task.state == 'FAILURE':
        return JsonResponse({'status': 'Failed'})

    else:
        # Task is still in progress
        return JsonResponse({'status': 'In Progress'})