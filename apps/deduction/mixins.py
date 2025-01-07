import io
import logging
import os
import re
from itertools import groupby

from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.staticfiles import finders
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.forms.models import model_to_dict
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import (
    CreateView,
    DeleteView,
    DetailView,
    ListView,
    UpdateView,
    View,
)
from docx import Document

from excalibur.rabbitmq_service import send_rabbitmq_message

from ..facture.parsing_docx import replace_placeholders_in_doc, set_font_size
from .constants import DEDUCTION_FIELDS
from .models import Deduction

logger = logging.getLogger(__name__)


class CheckOwnerDeduction:
    def __init__(self, model, request, kwargs, pk_url_kwarg='pk'):
        self.model = model
        self.request = request
        self.kwargs = kwargs
        self.pk_url_kwarg = pk_url_kwarg

    def get_queryset(self):
        """
        If user is superuser he can delete concrete deduction
        """
        if self.request.user.is_superuser:
            return self.model.objects.all()
        return self.model.objects.filter(owner=self.request.user)

    def get_deduction(self, queryset=None):
        """
        Ensure that only objects within the restricted queryset can be accessed.
        """
        queryset = self.get_queryset() if queryset is None else queryset
        pk = self.kwargs.get(self.pk_url_kwarg)

        try:
            obj = queryset.get(pk=pk)
        except self.model.DoesNotExist:
            raise PermissionDenied('you are not an owner of this deduction!')

        return obj


class BaseDeductionListView(LoginRequiredMixin, ListView):
    context_object_name = 'deductions'
    paginate_by = 3

    def get_queryset(self):
        supplier = self.request.GET.get('supplier')
        logger.error(f"Supplier filter: {supplier}")
        queryset = super().get_queryset().select_related('owner')

        if supplier:
            queryset = queryset.filter(supplier__icontains=supplier)
            logger.debug(f"Filtered queryset: {queryset}")
        else:
            queryset = queryset.order_by('-update_time')

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        supplier = self.request.GET.get('supplier', '')
        deductions = Deduction.objects.select_related('owner')
        if supplier:
            deductions = deductions.filter(supplier__icontains=supplier)
        deductions = deductions.order_by('supplier')

        grouped_deductions = {
            supplier: list(records)
            for supplier, records in groupby(deductions, key=lambda d: d.supplier)
        }

        logger.debug(f"Grouped deductions: {grouped_deductions}")

        grouped_items = list(grouped_deductions.items())
        paginator = Paginator(grouped_items, self.paginate_by)
        page_number = self.request.GET.get('page')
        page_obj = paginator.get_page(page_number)

        context['suppliers_with_deductions'] = page_obj
        context['page_obj'] = page_obj
        context['supplier_filter'] = supplier

        logger.debug(f"Context data: {context}")
        return context


class BaseDeductionCreateView(LoginRequiredMixin, CreateView):
    success_url = reverse_lazy('success_url')

    def form_valid(self, form):
        if not form.instance.owner:
            form.instance.owner = self.request.user
        response = super().form_valid(form)

        # Send RabbitMQ message on successful creation
        message = f"Deduction created: ID={self.object.id}, Owner={self.object.owner}"
        send_rabbitmq_message(self.model, message)

        return response


class BaseDeductionUpdateView(LoginRequiredMixin, UpdateView):
    def get_object(self, queryset=None):
        return CheckOwnerDeduction(self.model, self.request, self.kwargs).get_deduction()

    def form_valid(self, form):
        if not form.instance.owner:
            form.instance.owner = self.request.user
        response = super().form_valid(form)

        # Send RabbitMQ message on successful update
        message = f"Deduction updated: ID={self.object.id}, Owner={self.object.owner}"
        send_rabbitmq_message(self.model, message)

        return response


class BaseDeductionDeleteView(LoginRequiredMixin, DeleteView):
    def get_object(self, queryset=None):
        return CheckOwnerDeduction(self.model, self.request, self.kwargs).get_deduction()

    def delete(self, request, *args, **kwargs):
        obj = self.get_object()
        response = super().delete(request, *args, **kwargs)

        # Send RabbitMQ message on successful deletion
        message = f"Deduction deleted: ID={obj.id}, Owner={obj.owner}"
        send_rabbitmq_message(self.model, message)

        return response


class BaseDeductionDetailView(LoginRequiredMixin, DetailView):
    context_object_name = 'deduction'
    slug_url_kwarg = 'pk'

    def get_object(self, **kwargs):
        return get_object_or_404(self.model, pk=self.kwargs.get(self.slug_url_kwarg))


class BaseDeductionExportDocx(LoginRequiredMixin, View):
    local_template = 'deduction/file_templates/file_input/Deduction_template.docx'

    def generate_docx(self, facture_object, template_path):
        facture_dict = model_to_dict(facture_object)
        template_path = finders.find(template_path)
        doc = Document(template_path)

        for field in DEDUCTION_FIELDS:
            regex = re.compile(rf'{re.escape(field)}')
            replace_placeholders_in_doc(doc, regex, facture_dict)

        set_font_size(doc)
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream

    def get(self, request, *args, **kwargs):
        deduction_object = get_object_or_404(Deduction, pk=kwargs.get('pk'))
        file_stream = self.generate_docx(deduction_object, self.local_template)

        file_path = f'/usr/src/app/apps/deduction/static/deduction/file_templates/file_output/deduction_{deduction_object.number_deduction}.docx'
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, 'wb') as docx_file:
            docx_file.write(file_stream.getvalue())

        response = HttpResponse(
            file_stream,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="deduction_{deduction_object.number_deduction}.docx"'
        )

        # Send RabbitMQ message on document export
        message = f"Deduction document exported: ID={deduction_object.id}, Owner={deduction_object.owner}"
        send_rabbitmq_message(Deduction, message)

        return response
